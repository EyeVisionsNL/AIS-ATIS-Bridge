from __future__ import annotations
from io import BytesIO
from typing import Any
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADERS=("Mode","Channel Bank","ID","Name","Channel MHz","Frequency MHz","Scan Enabled")

def export_channels(settings:dict[str,Any])->bytes:
    workbook=Workbook(); sheet=workbook.active; sheet.title="Channels"; sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.fill=PatternFill("solid",fgColor="173746"); cell.font=Font(color="FFFFFF",bold=True); cell.alignment=Alignment(horizontal="center")
    for channel in settings["channels"]: sheet.append(("Marine",settings["channel_bank"],channel["id"],channel["label"],None,channel["frequency_mhz"],channel["scan_enabled"]))
    sheet.freeze_panes="A2"; sheet.auto_filter.ref=f"A1:G{sheet.max_row}"
    for width,column in zip((14,22,22,30,16,18,16),"ABCDEFG"): sheet.column_dimensions[column].width=width
    for cell in sheet["F"][1:]: cell.number_format="0.000000"
    output=BytesIO(); workbook.save(output); return output.getvalue()

def import_channels(data:bytes)->dict[str,Any]:
    if len(data)>2*1024*1024: raise ValueError("Excel file exceeds 2 MB")
    try: workbook=load_workbook(BytesIO(data),read_only=True,data_only=True)
    except Exception as error: raise ValueError("File is not a readable .xlsx workbook") from error
    if "Channels" not in workbook.sheetnames: raise ValueError("Workbook needs a Channels sheet")
    sheet=workbook["Channels"]; header=tuple(str(v or "").strip() for v in next(sheet.iter_rows(min_row=1,max_row=1,values_only=True)))
    if header[:7]!=HEADERS: raise ValueError("Channels headers do not match the export format")
    channels=[]; bank=None
    for row_number,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),start=2):
        values=list(row)+[None]*max(0,7-len(row))
        if not any(v not in (None,"") for v in values[:7]): continue
        if str(values[0] or "").strip().lower()!="marine": raise ValueError(f"Row {row_number}: Mode must be Marine")
        row_bank=str(values[1] or "").strip() or "marine"
        if bank is not None and row_bank!=bank: raise ValueError(f"Row {row_number}: Channel Bank must be identical")
        bank=row_bank; enabled=values[6] if isinstance(values[6],bool) else str(values[6] or "").strip().lower() in ("1","true","yes","ja","on")
        channels.append({"id":str(values[2] or "").strip(),"label":str(values[3] or "").strip(),"frequency_mhz":values[5],"scan_enabled":enabled})
    if not channels: raise ValueError("Workbook contains no channels")
    return {"channel_bank":bank or "marine","channels":channels}
